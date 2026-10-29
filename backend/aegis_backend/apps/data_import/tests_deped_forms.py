"""
Tests for the DepEd SF2/SF9 template readers (``deped_forms``).

The readers exist so a filled DepEd form imports as-is: SF9's learner rows sit
below a learning-area band on a named sheet, and SF2 packs a month into a grid
with its school header on a separate ``COVER`` sheet.  Neither is a flat table,
so the generic row-1-header path cannot read them.

These tests pin the contract the readers depend on:

- recognition is by sheet name, and a plain flat workbook is *not* a form;
- the filled workbooks the generator ships read back row-for-row;
- the readers are structural, so a moved band or a missing sheet fails loudly
  instead of attaching a score to the wrong learner or the wrong area;
- ``.xlsm`` is accepted and legacy ``.xls`` is still rejected.

Positive cases read the committed workbooks, so they also prove the generator
and the readers agree.
"""

import io
from pathlib import Path

import openpyxl
from django.test import SimpleTestCase

from apps.data_import import deped_forms
from apps.data_import.constants import ImportErrorCode
from apps.data_import.xlsx_pipeline import XlsxValidationError, validate_raw_xlsx

DATA = Path(__file__).resolve().parents[4] / "docs" / "test-data"

SF9_TEMPLATE = "G4-10  School Form 9 (SY 26-27).xlsx"
SF2_TEMPLATE = "SF2 JUNIOR HS.xlsm"
SF9_WORKBOOK = "grades-sf9-bonifacio.xlsx"
SF9_EARLIER_WORKBOOK = "grades-sf9-terms-1-2-bonifacio.xlsx"
SF2_WORKBOOK = "attendance-sf2-bonifacio.xlsm"
FLAT_WORKBOOK = "grades-import.xlsx"

_AREA_LABELS = [label for label, _ in deped_forms.SF9_LABEL_TO_BASE]


def _open(name):
    """Load a shipped workbook the way the pipeline does."""
    return openpyxl.load_workbook(DATA / name, read_only=True, data_only=True)


def _validate(name):
    """Run a shipped workbook through the public XLSX validation entry point."""
    return validate_raw_xlsx(io.BytesIO((DATA / name).read_bytes()), name)


def _filled_band(worksheet):
    """Put the learning-area band where the reader looks for it."""
    for offset, label in enumerate(_AREA_LABELS):
        worksheet.cell(
            row=deped_forms.SF9_AREA_BAND_ROW,
            column=deped_forms.SF9_AREA_FIRST_COL + offset,
        ).value = label


def _sf9_skeleton():
    """
    A minimal workbook shaped like SF9, for cases the shipped files cannot make.

    A real form has the learning-area band on every ``TERM n SUMMARY`` sheet, so
    the skeleton does too; the caller fills in learners and scores.
    """
    workbook = openpyxl.Workbook()
    workbook.active.title = "FINAL GRADES"
    for sheet_name, _period in deped_forms.SF9_TERM_SHEETS:
        worksheet = workbook.create_sheet(sheet_name)
        _filled_band(worksheet)
    return workbook


class DepedFormRecognitionTest(SimpleTestCase):
    def test_sf9_template_is_recognised(self):
        workbook = _open(SF9_TEMPLATE)
        try:
            self.assertEqual(deped_forms.detect(workbook.sheetnames), "sf9")
        finally:
            workbook.close()

    def test_sf2_template_is_recognised(self):
        workbook = _open(SF2_TEMPLATE)
        try:
            self.assertEqual(deped_forms.detect(workbook.sheetnames), "sf2")
        finally:
            workbook.close()

    def test_flat_workbook_is_not_a_deped_form(self):
        workbook = _open(FLAT_WORKBOOK)
        try:
            self.assertIsNone(deped_forms.detect(workbook.sheetnames))
            with self.assertRaises(deped_forms.DepedFormError) as caught:
                deped_forms.read(workbook)
        finally:
            workbook.close()
        self.assertEqual(caught.exception.code, "INVALID_WORKBOOK")

    def test_flat_workbook_still_uses_the_generic_path(self):
        """Detection must not divert a non-DepEd workbook."""
        result = _validate(FLAT_WORKBOOK)
        self.assertIn("student_id", result.headers)
        self.assertGreater(result.raw_row_count, 0)

    def test_xlsm_extension_is_accepted(self):
        result = _validate(SF2_WORKBOOK)
        self.assertEqual(result.headers, deped_forms.ATTENDANCE_HEADERS)
        self.assertEqual(result.raw_row_count, len(result.rows))

    def test_legacy_xls_is_still_rejected(self):
        content = io.BytesIO((DATA / SF2_WORKBOOK).read_bytes())
        with self.assertRaises(XlsxValidationError) as caught:
            validate_raw_xlsx(content, "attendance.xls")
        self.assertEqual(
            caught.exception.code, ImportErrorCode.INVALID_FILE_EXTENSION
        )


class Sf9ReaderTest(SimpleTestCase):
    def test_reads_one_row_per_learner_and_learning_area(self):
        workbook = _open(SF9_WORKBOOK)
        try:
            headers, rows = deped_forms.read_sf9(workbook)
        finally:
            workbook.close()

        self.assertEqual(headers, deped_forms.GRADE_HEADERS)
        self.assertEqual(len(rows), 200, "25 learners x 8 learning areas")
        self.assertEqual({row["grading_period"] for row in rows}, {"Term 3"})
        self.assertEqual({row["section"] for row in rows}, {"Grade 8 - Bonifacio"})
        self.assertEqual(len({row["student_id"] for row in rows}), 25)
        self.assertEqual(
            {row["subject_code"] for row in rows},
            {
                "FIL-8",
                "ENG-8",
                "MATH-8",
                "SCI-8",
                "AP-8",
                "ESP-8",
                "TLE-8",
                "MAPEH-8",
            },
        )

    def test_earlier_terms_workbook_covers_two_periods(self):
        workbook = _open(SF9_EARLIER_WORKBOOK)
        try:
            _headers, rows = deped_forms.read_sf9(workbook)
        finally:
            workbook.close()

        self.assertEqual(len(rows), 400, "25 learners x 8 areas x 2 terms")
        self.assertEqual(
            {row["grading_period"] for row in rows}, {"Term 1", "Term 2"}
        )

    def test_template_with_no_filled_learners_yields_no_rows(self):
        """An unfilled form must import zero rows, not blanks."""
        workbook = _open(SF9_TEMPLATE)
        try:
            _headers, rows = deped_forms.read_sf9(workbook)
        finally:
            workbook.close()
        self.assertEqual(rows, [])

    def test_placeholder_grade_level_is_not_a_section(self):
        """
        The unfilled template caches ``0`` in GRADE LEVEL/SECTION.  Such a sheet
        must be skipped, not become "Grade 0 - 0" with real scores attached.
        """
        workbook = _sf9_skeleton()
        worksheet = workbook["TERM 3 SUMMARY"]
        worksheet.cell(row=9, column=7).value = 0
        worksheet.cell(row=10, column=7).value = 0
        worksheet.cell(row=15, column=3).value = "AEGIS-QA-S0026"
        worksheet.cell(row=15, column=11).value = 88

        _headers, rows = deped_forms.read_sf9(workbook)

        self.assertEqual(rows, [])

    def test_mapeh_sub_scores_are_averaged(self):
        """MAPEH is the one area SF9 splits into two columns."""
        workbook = _sf9_skeleton()
        worksheet = workbook["TERM 3 SUMMARY"]
        worksheet.cell(row=9, column=7).value = 8
        worksheet.cell(row=10, column=7).value = "Bonifacio"
        worksheet.cell(row=15, column=3).value = "AEGIS-QA-S0026"
        worksheet.cell(row=15, column=18).value = 80  # M A
        worksheet.cell(row=15, column=19).value = 90  # P E H

        _headers, rows = deped_forms.read_sf9(workbook)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["subject_code"], "MAPEH-8")
        self.assertEqual(rows[0]["score"], "85")

    def test_malformed_learning_area_band_is_rejected(self):
        workbook = _sf9_skeleton()
        workbook["TERM 2 SUMMARY"].cell(
            row=deped_forms.SF9_AREA_BAND_ROW,
            column=deped_forms.SF9_AREA_FIRST_COL,
        ).value = "Filipino?"

        with self.assertRaises(deped_forms.DepedFormError) as caught:
            deped_forms.read_sf9(workbook)

        self.assertEqual(caught.exception.code, "MALFORMED_HEADER")
        self.assertIn("TERM 2 SUMMARY", caught.exception.message)


class Sf2ReaderTest(SimpleTestCase):
    def test_reads_one_row_per_learner_and_school_day(self):
        workbook = _open(SF2_WORKBOOK)
        try:
            headers, rows = deped_forms.read_sf2(workbook)
        finally:
            workbook.close()

        self.assertEqual(headers, deped_forms.ATTENDANCE_HEADERS)
        self.assertEqual(len(rows), 125, "25 learners x 5 school days")
        self.assertEqual(len({row["student_id"] for row in rows}), 25)
        self.assertEqual({row["section"] for row in rows}, {"Grade 8 - Bonifacio"})
        self.assertEqual(
            sorted({row["date"] for row in rows}),
            [
                "2026-09-07",
                "2026-09-08",
                "2026-09-09",
                "2026-09-10",
                "2026-09-11",
            ],
        )
        self.assertTrue(
            {row["status"] for row in rows}
            <= {deped_forms.SF2_DEFAULT_STATUS, "late", "absent"}
        )

    def test_blank_cells_read_as_present(self):
        """SF2 marks the exceptions only: a blank day is a present day."""
        workbook = _open(SF2_WORKBOOK)
        try:
            _headers, rows = deped_forms.read_sf2(workbook)
        finally:
            workbook.close()
        self.assertGreater(
            sum(1 for row in rows if row["status"] == "present"),
            0,
            "the shipped week has no present days",
        )

    def test_missing_cover_sheet_is_rejected(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = deped_forms.SF2_SHEET

        with self.assertRaises(deped_forms.DepedFormError) as caught:
            deped_forms.read_sf2(workbook)

        self.assertEqual(caught.exception.code, "MISSING_REQUIRED_COLUMN")
